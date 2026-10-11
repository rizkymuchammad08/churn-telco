"""P6: the dashboard runs without exceptions; shows model_hash, the contact list, and the synthetic banner.

The fixtures build a mini project in tmp_path through the real CLI (train + roi), so the artifacts
the dashboard reads are exactly what the pipeline produces and the repo is left untouched.
"""

import json
import shutil
import subprocess
import sys

import pytest
import yaml

from src.config import ROOT
from src.data.validate import BANNER
from src.models.predict import predict
from src.models.train import load_bundle
from src.reporting.dashboard import DashboardError, contact_table, load_view

pytestmark = pytest.mark.p6

# The dashboard is an optional extra: without streamlit and openpyxl these tests are skipped, not errors.
pytest.importorskip("openpyxl")
AppTest = pytest.importorskip("streamlit.testing.v1").AppTest

APP = ROOT / "src" / "reporting" / "dashboard.py"
TEXT_KINDS = (
    "markdown",
    "text",
    "code",
    "error",
    "warning",
    "info",
    "success",
    "caption",
    "title",
    "subheader",
)


# Words from the dashboard UI; change them here if the UI text changes.
OF_WORD = "of"  # the UI reads "<n> of <k> customers"
CONDITIONAL_WORD = "conditional"


def _run_cli(root, *module_args):
    r = subprocess.run(
        [sys.executable, "-m", *module_args], cwd=root, capture_output=True, text=True, timeout=300
    )
    assert r.returncode == 0, r.stdout + r.stderr
    return r.stdout


@pytest.fixture(scope="module")
def project(tmp_path_factory):
    root = tmp_path_factory.mktemp("proj")
    shutil.copytree(ROOT / "src", root / "src", ignore=shutil.ignore_patterns("__pycache__"))
    for f in ("params.yaml", "assumptions.yaml"):
        shutil.copy(ROOT / f, root / f)
    (root / "data").mkdir()
    shutil.copy(ROOT / "data" / "sample.csv", root / "data" / "sample.csv")
    _run_cli(root, "src.models.train", "--sample")
    _run_cli(root, "src.reporting.roi")
    return root


@pytest.fixture(scope="module")
def synthetic_project(project, tmp_path_factory):
    root = tmp_path_factory.mktemp("synth")
    shutil.copytree(project, root, dirs_exist_ok=True)
    p = yaml.safe_load((root / "params.yaml").read_text(encoding="utf-8"))
    p["data_source"] = "synthetic"
    (root / "params.yaml").write_text(yaml.safe_dump(p, sort_keys=False), encoding="utf-8")
    mpath = root / "reports" / "metrics.json"
    m = json.loads(mpath.read_text(encoding="utf-8"))
    m["data_source"] = "synthetic"
    mpath.write_text(json.dumps(m), encoding="utf-8")
    return root


def run_app(root):
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("CHURN_ROOT", str(root))
        return AppTest.from_file(str(APP), default_timeout=120).run()


def texts(at):
    out = []
    for kind in TEXT_KINDS:
        out += [str(el.value) for el in getattr(at, kind)]
    return out


@pytest.fixture(scope="module")
def app(project):
    return run_app(project)


@pytest.fixture(scope="module")
def synth_app(synthetic_project):
    return run_app(synthetic_project)


@pytest.fixture(scope="module")
def metrics(project):
    return json.loads((project / "reports" / "metrics.json").read_text(encoding="utf-8"))


# --- runs without exception ------------------------------------------------------
def test_app_runs_without_exception(app, synth_app):
    assert list(app.exception) == []
    assert list(synth_app.exception) == []
    assert list(app.error) == []  # real data: no banner and no error


# --- model_hash --------------------------------------------------------------------
def test_app_shows_model_hash_from_metrics(app, metrics):
    h = metrics["hashes"]["model"]
    assert len(h) == 64
    assert h in [str(c.value) for c in app.code]


def test_model_hash_matches_bundle_saved_by_cli_in_a_fresh_process(project, metrics):
    # The bundle from `python -m src.models.train` must be loadable by another process
    # (not a pickled __main__.* object).
    code = (
        "from src.models.train import load_bundle;"
        "print(load_bundle('artifacts/model.joblib').model.model_hash)"
    )
    r = subprocess.run([sys.executable, "-c", code], cwd=project, capture_output=True, text=True)
    assert r.returncode == 0, r.stderr
    assert r.stdout.strip() == metrics["hashes"]["model"]


# --- contact list ------------------------------------------------------------------
def _expected_contacts(root):
    """Independent computation: top-k by rank, then p_churn >= p* (without using the targeting module)."""
    a = yaml.safe_load((root / "assumptions.yaml").read_text(encoding="utf-8"))
    params = yaml.safe_load((root / "params.yaml").read_text(encoding="utf-8"))
    from src.data.load import load_raw

    df = load_raw(root / params["paths"]["sample"], params)
    pred = predict(df, load_bundle(root / "artifacts" / "model.joblib"), params)
    n = len(pred)
    k = min(n, int(n * a["campaign_capacity"] / a["population_per_period"] + 0.5))
    p_star = a["cost_intervention"] / (a["save_rate"] * a["clv"])
    top = pred.sort_values("rank").head(k)
    return list(top[top["p_churn"] >= p_star].index), k, p_star


def test_app_shows_contact_list(app, project):
    ids, k, p_star = _expected_contacts(project)
    assert 0 < len(ids) <= k
    tables = [d.value for d in app.dataframe]
    contacts = next(t for t in tables if "customerID" in t.columns)
    assert list(contacts["customerID"]) == ids
    assert list(contacts["rank"]) == sorted(contacts["rank"])
    assert (contacts["p_churn"] >= p_star).all()
    assert contacts["customerID"].is_unique
    assert any(f"{len(ids)} {OF_WORD}" in s for s in texts(app))


def test_view_and_table_agree(project):
    v = load_view(project)
    t = contact_table(v)
    assert list(t.columns) == ["customerID", "rank", "p_churn", "out_of_domain"]
    assert len(t) == v.plan.n_contacts <= v.k
    assert v.model_hash == json.loads((project / "reports/metrics.json").read_text())["hashes"]["model"]


# --- synthetic banner ----------------------------------------------------------------
def test_banner_shown_for_synthetic(synth_app):
    assert BANNER in [str(e.value) for e in synth_app.error]


def test_banner_absent_for_real_data(app):
    assert all(BANNER not in s for s in texts(app))


def test_synthetic_app_still_shows_hash_and_contacts(synth_app, metrics):
    assert metrics["hashes"]["model"] in [str(c.value) for c in synth_app.code]
    assert any("customerID" in d.value.columns for d in synth_app.dataframe)


def test_conditional_value_warning_while_save_rate_assumed(app):
    assert any(CONDITIONAL_WORD in str(w.value).lower() for w in app.warning)


# --- artifact errors are displayed, not raised as exceptions ----------------------------
def test_missing_artifacts_shown_as_error_not_exception(project, tmp_path):
    root = tmp_path / "bare"
    shutil.copytree(project, root, ignore=shutil.ignore_patterns("artifacts"))
    at = run_app(root)
    assert list(at.exception) == []
    assert any("E_DASH_MISSING" in str(e.value) for e in at.error)
    with pytest.raises(DashboardError) as e:
        load_view(root)
    assert e.value.code == "E_DASH_MISSING"


def test_stale_model_hash_shown_as_error(project, tmp_path):
    root = tmp_path / "stale"
    shutil.copytree(project, root)
    mpath = root / "reports" / "metrics.json"
    m = json.loads(mpath.read_text(encoding="utf-8"))
    m["hashes"]["model"] = "0" * 64
    mpath.write_text(json.dumps(m), encoding="utf-8")
    at = run_app(root)
    assert list(at.exception) == []
    assert any("E_DASH_STALE" in str(e.value) for e in at.error)


# --- upload and download -------------------------------------------------------------
def _sample_bytes(root, n=300):
    import pandas as pd

    df = pd.read_csv(root / "data" / "sample.csv", dtype=str, keep_default_na=False).head(n)
    return df, df.to_csv(index=False).encode("utf-8")


def _scoring(root):
    from src.reporting.dashboard import load_scoring

    return load_scoring(root)


def test_upload_csv_scores_like_the_pipeline(project):
    from src.reporting.dashboard import plan_table, read_upload, score_upload

    bundle, params, assumptions = _scoring(project)
    df, raw = _sample_bytes(project)
    plan = score_upload(read_upload("x.csv", raw, params), bundle, params, assumptions)
    table = plan_table(plan)
    assert list(table.columns) == ["customerID", "rank", "p_churn", "out_of_domain"]
    assert table["rank"].is_monotonic_increasing
    assert (table["p_churn"] >= plan.p_star).all()
    assert set(table["customerID"]) <= set(df["customerID"])


def test_upload_xlsx_equals_csv_and_label_column_is_optional(project):
    import io

    from src.reporting.dashboard import plan_table, read_upload, score_upload, to_xlsx_bytes

    bundle, params, assumptions = _scoring(project)
    df, raw = _sample_bytes(project)
    buf = io.BytesIO()
    df.drop(columns=["Churn"]).to_excel(buf, index=False, engine="openpyxl")
    from_csv = plan_table(score_upload(read_upload("x.csv", raw, params), bundle, params, assumptions))
    from_xlsx = plan_table(
        score_upload(read_upload("x.xlsx", buf.getvalue(), params), bundle, params, assumptions)
    )
    assert from_csv["customerID"].tolist() == from_xlsx["customerID"].tolist()
    assert to_xlsx_bytes(from_xlsx)[:2] == b"PK"  # a valid xlsx is a zip file


def test_download_csv_round_trips(project):
    import io

    import pandas as pd

    from src.reporting.dashboard import plan_table, read_upload, score_upload, to_csv_bytes

    bundle, params, assumptions = _scoring(project)
    _, raw = _sample_bytes(project)
    table = plan_table(score_upload(read_upload("x.csv", raw, params), bundle, params, assumptions))
    back = pd.read_csv(io.BytesIO(to_csv_bytes(table)))
    assert back["customerID"].tolist() == table["customerID"].tolist()
    assert len(back) == len(table)


def test_upload_rejects_bad_files_with_stable_codes(project):
    from src.reporting.dashboard import read_upload, score_upload

    bundle, params, assumptions = _scoring(project)
    with pytest.raises(DashboardError) as e:
        read_upload("x.txt", b"a", params)
    assert e.value.code == "E_UPLOAD_TYPE"
    with pytest.raises(DashboardError) as e:
        read_upload("x.xlsx", b"not a workbook", params)
    assert e.value.code == "E_UPLOAD_READ"
    with pytest.raises(DashboardError) as e:
        read_upload("x.csv", b"", params)
    assert e.value.code == "E_UPLOAD_READ"
    from src.models.predict import PredictError

    df = read_upload("x.csv", b"customerID,tenure\nA,1\n", params)
    with pytest.raises(PredictError) as e:
        score_upload(df, bundle, params, assumptions)
    assert e.value.code == "E_PREDICT_COLUMNS"


# --- Export: customerID from an upload must never become a spreadsheet formula ------------------

FORMULA_IDS = ['=HYPERLINK("http://example.invalid","x")', "+1+1", "-2+3", "@SUM(1)", "\t=1+1"]
SAFE_ID = "7590-VHVEG"


def _formula_table():
    import pandas as pd

    ids = [*FORMULA_IDS, SAFE_ID]
    n = len(ids)
    return pd.DataFrame(
        {"customerID": ids, "rank": range(1, n + 1), "p_churn": [0.9] * n, "out_of_domain": [False] * n}
    )


def test_csv_export_neutralizes_formula_cells():
    import csv
    import io

    from src.reporting.dashboard import to_csv_bytes

    rows = list(csv.reader(io.StringIO(to_csv_bytes(_formula_table()).decode("utf-8"), newline="")))
    got = [r[0] for r in rows[1:]]
    assert got == ["'" + x for x in FORMULA_IDS] + [SAFE_ID]


def test_xlsx_export_stores_formula_like_ids_as_text():
    import io

    from openpyxl import load_workbook

    from src.reporting.dashboard import to_xlsx_bytes

    ws = load_workbook(io.BytesIO(to_xlsx_bytes(_formula_table()))).active
    cells = [ws.cell(row=r, column=1) for r in range(2, ws.max_row + 1)]
    assert all(c.data_type == "s" for c in cells)  # none is a formula
    assert [c.value for c in cells] == [*FORMULA_IDS, SAFE_ID]  # values are kept exactly
